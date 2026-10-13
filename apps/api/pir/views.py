from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
import uuid
from django.contrib.auth import authenticate, login, logout, get_user_model
from django.core.cache import cache
from django.db import transaction, IntegrityError
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.views.decorators.csrf import ensure_csrf_cookie, csrf_protect
from django.views.decorators.http import require_POST
from rest_framework.decorators import api_view
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.response import Response
from rest_framework.views import exception_handler as drf_exception_handler
from .models import Project, Schedule, Revision, Fact, Audit, Calendar, NetworkTemplate, ReportSnapshot
from .planning import DomainError, calculate, day, add_workdays, money, working_days
from .services import role, accessible, serialize, cutoff
from .exporting import workbook


def exception_handler(exc, context):
    if isinstance(exc, (DomainError, IntegrityError)):
        return Response({'detail': str(exc) if isinstance(exc, DomainError) else 'Такой шифр или номер уже существует.'}, status=400)
    return drf_exception_handler(exc, context)


def allow(user, *roles):
    if role(user) not in roles:
        raise PermissionDenied('Недостаточно прав для этого действия.')


def person(user):
    return {'id': user.pk, 'name': user.get_full_name() or user.username, 'username': user.username, 'role': role(user)}


@ensure_csrf_cookie
def session(request):
    return JsonResponse({'user': person(request.user) if request.user.is_authenticated else None, 'today': timezone.localdate().isoformat()})


@require_POST
@csrf_protect
def sign_in(request):
    import json
    try: data = json.loads(request.body)
    except (ValueError, TypeError): return JsonResponse({'detail': 'Некорректный запрос'}, status=400)
    if not isinstance(data, dict): return JsonResponse({'detail': 'Некорректный запрос'}, status=400)
    username = str(data.get('username', ''))[:150]
    key = 'login:'+request.META.get('REMOTE_ADDR', '')+':'+username
    if cache.get(key, 0) >= 10: return JsonResponse({'detail': 'Слишком много попыток. Повторите через 10 минут.'}, status=429)
    user = authenticate(request, username=username, password=data.get('password', ''))
    if not user:
        cache.set(key, cache.get(key, 0)+1, 600)
        return JsonResponse({'detail': 'Неверный логин или пароль.'}, status=400)
    cache.delete(key); login(request, user)
    return JsonResponse({'user': person(user)})


@api_view(['POST'])
def sign_out(request):
    logout(request)
    return Response({'ok': True})


@api_view(['GET'])
def bootstrap(request):
    users = get_user_model().objects.filter(profile__role__in=['gip', 'technical', 'finance'])
    projects = Project.objects.all()
    if role(request.user) == 'gip': projects = projects.filter(gip=request.user)
    if role(request.user) == 'department': projects = projects.filter(schedules__in=accessible(request.user)).distinct()
    return Response({'user': person(request.user), 'today': timezone.localdate().isoformat(), 'users': [person(u) for u in users], 'projects': [{'id': str(p.pk), 'code': p.code, 'title': p.title, 'gip_id': p.gip_id} for p in projects], 'templates': [{'id': t.pk, 'code': t.code, 'title': t.title, 'is_demo': t.is_demo, 'tasks': t.data.get('tasks', [])} for t in NetworkTemplate.objects.all()], 'calendars': [{'id': c.pk, 'title': c.title, 'is_demo': c.is_demo, 'data': c.data} for c in Calendar.objects.all()]})


@api_view(['GET', 'POST'])
def schedules(request):
    if request.method == 'GET':
        return Response([serialize(s, request.user) for s in accessible(request.user).order_by('-created_at')])
    allow(request.user, 'gip')
    d = request.data
    with transaction.atomic():
        if d.get('project_id'):
            project = get_object_or_404(Project, pk=d['project_id'], gip=request.user)
        else:
            if not str(d.get('project_code', '')).strip() or not str(d.get('project_title', '')).strip():
                raise DomainError('Укажите шифр и наименование объекта.')
            project = Project.objects.create(code=d['project_code'], title=d['project_title'], gip=request.user)
        if not d.get('number') or not d.get('title'): raise DomainError('Укажите номер и наименование графика.')
        calendar = get_object_or_404(Calendar, pk=d.get('calendar_id'))
        tasks = []
        template = NetworkTemplate.objects.filter(pk=d.get('template_id')).first() if d.get('template_id') else None
        start = day(d.get('start')).isoformat()
        if template:
            for item in template.data.get('tasks', []):
                t = deepcopy(item)
                t['start'] = add_workdays(start, t.pop('offset', 0), calendar.data)
                duration = t.pop('duration', 1)
                t['end'] = add_workdays(t['start'], max(0, duration-1), calendar.data)
                t['cost'] = '0'
                tasks.append(t)
        end = max([t['end'] for t in tasks]+[start])
        data = {'start': start, 'end': end, 'budget': '0', 'calendar': {**calendar.data, 'title': calendar.title, 'is_demo': calendar.is_demo, 'id': calendar.pk}, 'model_code': template.code if template else 'Индивидуальная', 'tasks': tasks}
        calculate(data)
        s = Schedule.objects.create(project=project, number=d['number'], title=d['title'], stage=d.get('stage', 'РД'), phase=d.get('phase', '01'))
        Revision.objects.create(schedule=s, number=1, data=data, author=request.user)
        Audit.objects.create(schedule=s, user=request.user, action='График создан')
    return Response(serialize(s, request.user), status=201)


@api_view(['GET', 'PUT'])
def detail(request, pk):
    s = get_object_or_404(accessible(request.user), pk=pk)
    if request.method == 'GET':
        revision = get_object_or_404(Revision, schedule=s, pk=request.query_params['revision']) if request.query_params.get('revision') else None
        return Response(serialize(s, request.user, revision=revision))
    allow(request.user, 'gip')
    with transaction.atomic():
        Schedule.objects.select_for_update().get(pk=s.pk)
        v = s.revisions.order_by('-number').first()
        if v.status != 'draft': raise DomainError('Для изменения создайте новую редакцию.')
        if request.data.get('edit_version') != v.edit_version: return Response({'detail': 'График уже изменён. Обновите страницу.'}, status=412)
        payload = request.data.get('data')
        if not isinstance(payload, dict): raise DomainError('Некорректные данные графика.')
        # Calendar is an immutable selected server-side version, never a client-defined holiday list.
        cal = get_object_or_404(Calendar, pk=payload.get('calendar', {}).get('id'))
        payload['calendar'] = {**cal.data, 'title': cal.title, 'is_demo': cal.is_demo, 'id': cal.pk}
        calc = calculate(payload)
        payload['budget'] = str(money(payload['budget']))
        for task in payload['tasks']: task['cost'] = str(money(task.get('cost', '0')))
        v.data = payload; v.edit_version += 1; v.save(update_fields=['data', 'edit_version'])
        Audit.objects.create(schedule=s, user=request.user, action='Черновик изменён', detail={'revision': v.number, 'edit_version': v.edit_version})
    s = accessible(request.user).get(pk=pk)
    return Response(serialize(s, request.user))


@api_view(['POST'])
def action(request, pk, command):
    s = get_object_or_404(accessible(request.user), pk=pk)
    with transaction.atomic():
        Schedule.objects.select_for_update().get(pk=s.pk)
        v = s.revisions.order_by('-number').first()
        reason = str(request.data.get('reason', '')).strip()
        if command == 'revise':
            allow(request.user, 'gip')
            if v.status in ['draft', 'review']: raise DomainError('У графика уже есть рабочая редакция.')
            if not reason: raise DomainError('Укажите причину изменения.')
            v = Revision.objects.create(schedule=s, number=v.number+1, data=deepcopy(v.data), author=request.user, reason=reason)
        elif command == 'submit':
            allow(request.user, 'gip')
            if v.status != 'draft': raise DomainError('Отправить можно только черновик.')
            calc = calculate(v.data)
            if calc['errors']: raise DomainError('; '.join(calc['errors']))
            approver = get_object_or_404(get_user_model(), pk=request.data.get('approver_id'), profile__role__in=['technical', 'finance'])
            if approver.pk == request.user.pk: raise DomainError('Самоутверждение запрещено.')
            v.status = 'review'; v.approver = approver; v.submitted_at = timezone.now(); v.save()
        elif command in ['approve', 'return']:
            allow(request.user, 'technical', 'finance')
            if v.status != 'review' or v.approver_id != request.user.pk: raise DomainError('Эта редакция не назначена вам на утверждение.')
            if v.author_id == request.user.pk: raise DomainError('Самоутверждение запрещено.')
            if command == 'approve':
                calc = calculate(v.data)
                if calc['errors']: raise DomainError('; '.join(calc['errors']))
                v.status = 'approved'; v.approved_at = timezone.now()
            else:
                if not reason: raise DomainError('Укажите причину возврата.')
                v.status = 'returned'; v.reason = reason
            v.save()
        else: raise DomainError('Неизвестное действие.')
        Audit.objects.create(schedule=s, user=request.user, action={'revise': 'Создана новая редакция', 'submit': 'Отправлен на утверждение', 'approve': 'График утверждён', 'return': 'Возвращён на доработку'}[command], detail={'revision': v.number, 'reason': reason})
    return Response(serialize(accessible(request.user).get(pk=pk), request.user))


@api_view(['POST'])
def facts(request, pk):
    allow(request.user, 'coordinator')
    s = get_object_or_404(accessible(request.user), pk=pk)
    v = s.revisions.filter(status='approved').last()
    if not v: raise DomainError('Сначала утвердите график.')
    d = request.data
    task = next((t for t in v.data['tasks'] if t['id'] == d.get('node_id')), None)
    if not task: raise DomainError('Работа не найдена в действующем плане.')
    effective = day(d.get('effective_date'))
    progress = money(d.get('progress', 0))
    start = day(d['actual_start']) if d.get('actual_start') else None
    end = day(d['actual_end']) if d.get('actual_end') else None
    reason = str(d.get('reason', '')).strip()
    if progress > 100 or effective > timezone.localdate(): raise DomainError('Процент должен быть 0–100, дата состояния — не в будущем.')
    if progress > 0 and not start: raise DomainError('Для выполнения укажите фактическое начало.')
    if (progress == 100) != bool(end): raise DomainError('100% готовности и дата окончания указываются вместе.')
    if end and (not start or end < start): raise DomainError('Окончание не может быть раньше начала.')
    if (start and start > effective) or (end and end > effective): raise DomainError('Фактические даты не могут быть позже даты состояния.')
    previous = s.facts.filter(node_id=task['id']).order_by('-effective_date', '-recorded_at').first()
    if ((task['end'] < effective.isoformat() and progress < 100) or (previous and progress < previous.progress)) and not reason:
        raise DomainError('Укажите причину задержки или снижения готовности.')
    with transaction.atomic():
        Fact.objects.create(schedule=s, node_id=task['id'], effective_date=effective, progress=progress, actual_start=start, actual_end=end, reason=reason, author=request.user)
        Audit.objects.create(schedule=s, user=request.user, action='Обновлён факт', detail={'node_id': task['id'], 'progress': str(progress), 'effective_date': str(effective), 'reason': reason})
    return Response(serialize(accessible(request.user).get(pk=pk), request.user))


@api_view(['GET'])
def history(request, pk):
    s = get_object_or_404(accessible(request.user), pk=pk)
    return Response([{'id': a.pk, 'action': a.action, 'user': a.user.get_full_name(), 'date': a.created_at.isoformat(), 'detail': a.detail} for a in s.audit.select_related('user').all()[:200]])


def report_rows(request):
    as_of = request.query_params.get('as_of', timezone.localdate().isoformat())
    day(as_of)
    known = cutoff(request.query_params['known_at']) if request.query_params.get('known_at') else None
    return [row for s in accessible(request.user) if (row := serialize(s, request.user, as_of=as_of, known_at=known, historical=True))]


@api_view(['GET'])
def report(request):
    return Response(report_rows(request))


def file_response(data, name):
    response = HttpResponse(data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="'+name+'"'
    response['Cache-Control'] = 'no-store'
    return response


@api_view(['GET'])
def export(request, pk=None):
    rows = [serialize(get_object_or_404(accessible(request.user), pk=pk), request.user)] if pk else report_rows(request)
    Audit.objects.create(user=request.user, action='Выгрузка Excel', detail={'count': len(rows)})
    return file_response(workbook(rows), 'pir-schedule.xlsx' if pk else 'pir-report.xlsx')


@api_view(['GET', 'POST'])
def snapshots(request):
    if request.method == 'GET':
        return Response([{'id': str(s.pk), 'title': s.title, 'created_at': s.created_at.isoformat(), 'count': len(s.data)} for s in ReportSnapshot.objects.filter(owner=request.user).order_by('-created_at')[:100]])
    with transaction.atomic():
        from django.db import connection
        if connection.vendor == 'postgresql':
            with connection.cursor() as cursor: cursor.execute('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ')
        rows = report_rows(request)
        snap = ReportSnapshot.objects.create(owner=request.user, title='Планёрка · '+request.query_params.get('as_of', timezone.localdate().isoformat()), data=rows)
    return Response({'id': str(snap.pk)}, status=201)


@api_view(['GET'])
def snapshot_export(request, pk):
    snap = get_object_or_404(ReportSnapshot, pk=pk, owner=request.user)
    ids = set(str(s.pk) for s in accessible(request.user))
    if any(s['id'] not in ids for s in snap.data) or role(request.user) == 'department':
        raise PermissionDenied('Права на данные снимка изменились. Сформируйте новый отчёт.')
    return file_response(workbook(snap.data, snap.title, True), 'pir-snapshot.xlsx')


@api_view(['POST'])
def calendars(request):
    allow(request.user, 'coordinator', 'technical')
    d = request.data
    start, end = day(d.get('start')), day(d.get('end'))
    if end < start or (end-start).days > 3660: raise DomainError('Некорректный период календаря.')
    for value in d.get('holidays', [])+d.get('working_days', []):
        if not start <= day(value) <= end: raise DomainError('Исключение вне периода календаря.')
    if not d.get('title'): raise DomainError('Укажите название и версию календаря.')
    c = Calendar.objects.create(title=d['title'], data={'start': str(start), 'end': str(end), 'holidays': d.get('holidays', []), 'working_days': d.get('working_days', []), 'source': d.get('source', '')}, is_demo=not bool(d.get('verified')))
    Audit.objects.create(user=request.user, action='Создана версия календаря', detail={'calendar_id': c.pk})
    return Response({'id': c.pk}, status=201)


@api_view(['POST'])
def import_tasks(request, pk):
    """MVP: preview-only XLSX import into the schedule editor. Save uses the normal revision API."""
    allow(request.user, 'gip')
    s = get_object_or_404(accessible(request.user), pk=pk)
    if s.revisions.last().status != 'draft': raise DomainError('Импорт доступен в черновик.')
    f = request.FILES.get('file')
    if not f or f.size > 5*1024*1024: raise DomainError('Выберите XLSX размером до 5 МБ.')
    from zipfile import ZipFile, BadZipFile
    from openpyxl import load_workbook
    raw = f.read()
    try:
        with ZipFile(BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > 30*1024*1024: raise DomainError('Слишком большой распакованный файл.')
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ValueError, KeyError): raise DomainError('Файл должен быть корректным XLSX.')
    sheet = book.active
    rows = list(sheet.iter_rows(max_row=1002, max_col=8))
    if len(rows) > 1001: raise DomainError('В MVP допускается до 1000 строк.')
    expected = ['ID', 'Работа', 'Шифр', 'Отдел', 'Автор', 'Начало', 'Окончание', 'Сумма']
    if [c.value for c in rows[0]] != expected: raise DomainError('Колонки: '+', '.join(expected))
    tasks = []
    for n, row in enumerate(rows[1:], 2):
        if not any(c.value is not None for c in row): continue
        if any(c.data_type == 'f' for c in row): raise DomainError(f'Строка {n}: формулы не допускаются.')
        values = [c.value for c in row]
        if not isinstance(values[2], str): raise DomainError(f'Строка {n}: шифр специальности должен быть текстом, чтобы сохранить ведущие нули.')
        def dateval(v): return v.date().isoformat() if hasattr(v, 'date') else day(v).isoformat()
        tasks.append({'id': str(values[0] or uuid.uuid4()), 'title': str(values[1] or ''), 'specialty': str(values[2] or ''), 'department': str(values[3] or ''), 'author': str(values[4] or ''), 'start': dateval(values[5]), 'end': dateval(values[6]), 'cost': str(money(values[7] or 0)), 'dependencies': [], 'milestone': False})
    book.close()
    return Response({'tasks': tasks, 'warning': 'Импорт заменит работы черновика после нажатия «Сохранить». Зависимости добавляются вручную.'})
