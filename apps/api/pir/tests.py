from copy import deepcopy
from datetime import datetime
from io import BytesIO
from django.contrib.auth import get_user_model
from django.test import TransactionTestCase, SimpleTestCase, Client
from django.utils import timezone
from rest_framework.test import APIClient
from openpyxl import load_workbook, Workbook
from .models import Profile, Project, Calendar, Schedule, Revision, Fact
from .planning import calculate, DomainError, money, working_days


CAL = {'start': '2026-01-01', 'end': '2026-12-31', 'holidays': [], 'working_days': []}


def task(i, start, end, deps=None, cost='0', **extra):
    return {'id': i, 'title': i, 'specialty': '062', 'department': 'Путь', 'author': 'Автор', 'start': start, 'end': end, 'cost': cost, 'dependencies': deps or [], **extra}


def fs(i): return {'from': i, 'type': 'FS', 'lag': 0}


def network():
    return {'start': '2026-01-05', 'end': '2026-01-16', 'budget': '100', 'calendar': deepcopy(CAL), 'model_code': 'TEST', 'tasks': [task('a','2026-01-05','2026-01-07'),task('b','2026-01-08','2026-01-14',[fs('a')],'60'),task('c','2026-01-08','2026-01-09',[fs('a')],'40'),task('d','2026-01-15','2026-01-16',[fs('b'),fs('c')])]}


class PlanningTests(SimpleTestCase):
    def test_parallel_critical_path(self):
        r=calculate(network())
        self.assertEqual(r['length'],10);self.assertEqual(r['nodes']['c']['float'],3);self.assertEqual(r['errors'],[])
        self.assertEqual({i for i,n in r['nodes'].items() if n['critical']},{'a','b','d'})

    def test_cycle(self):
        n=network();n['tasks'][0]['dependencies']=[fs('d')]
        with self.assertRaisesRegex(DomainError,'цикл'):calculate(n)

    def test_negative_deadline_reserve(self):
        n=network();n['end']='2026-01-15'
        r=calculate(n);self.assertEqual(r['nodes']['a']['deadline_float'],-1);self.assertFalse(r['valid'])

    def test_weekends_and_exceptions(self):
        self.assertEqual(len(working_days('2026-01-09','2026-01-12',CAL)),2)
        self.assertEqual(len(working_days('2026-01-09','2026-01-12',{**CAL,'holidays':['2026-01-12']})),1)

    def test_ss_long_predecessor_determines_finish(self):
        n=network();n['tasks']=[task('a','2026-01-05','2026-01-16',cost='100'),task('b','2026-01-05','2026-01-05',[{'from':'a','type':'SS','lag':0}])]
        r=calculate(n);self.assertEqual(r['length'],10);self.assertEqual(r['nodes']['b']['float'],9)

    def test_ff_relation(self):
        n=network();n['tasks']=[task('a','2026-01-05','2026-01-07',cost='100'),task('b','2026-01-06','2026-01-07',[{'from':'a','type':'FF','lag':0}])]
        r=calculate(n);self.assertEqual(r['nodes']['b']['early'],1);self.assertEqual(r['length'],3)

    def test_event_start_vs_end(self):
        n=network();n['tasks']=[task('a','2026-01-05','2026-01-05',milestone=True,position='start'),task('b','2026-01-05','2026-01-05',[fs('a')],cost='100')]
        self.assertTrue(calculate(n)['valid'])
        n['tasks'][0]['position']='end'
        self.assertFalse(calculate(n)['valid'])

    def test_invalid_money(self):
        for value in ['NaN','Infinity','-1','0.001']:
            with self.assertRaises(DomainError):money(value)

    def test_duplicate_or_unknown_link(self):
        n=network();n['tasks'][1]['dependencies']=[fs('missing')]
        with self.assertRaises(DomainError):calculate(n)


class WorkflowTests(TransactionTestCase):
    def setUp(self):
        self.users={}
        for r in ['gip','coordinator','technical','finance','director','department']:
            u=get_user_model().objects.create_user(r,password='TestPassword2026!')
            Profile.objects.create(user=u,role=r,department='Путь');self.users[r]=u
        self.other=get_user_model().objects.create_user('other',password='TestPassword2026!');Profile.objects.create(user=self.other,role='gip')
        self.calendar=Calendar.objects.create(title='Test',data=CAL)
        self.project=Project.objects.create(code='TEST',title='Test',gip=self.users['gip'])
        self.s=Schedule.objects.create(project=self.project,number='TEST.01',title='Test')
        data=network();data['calendar']={**CAL,'id':self.calendar.pk}
        self.v=Revision.objects.create(schedule=self.s,number=1,data=data,author=self.users['gip'])
        self.client=APIClient();self.client.force_authenticate(self.users['gip'])
        self.path=f'/api/schedules/{self.s.pk}'

    def approve(self):
        self.client.force_authenticate(self.users['gip'])
        self.assertEqual(self.client.post(self.path+'/submit',{'approver_id':self.users['technical'].pk},format='json').status_code,200)
        self.client.force_authenticate(self.users['technical'])
        r=self.client.post(self.path+'/approve',{},format='json');self.assertEqual(r.status_code,200,r.data)
        self.v.refresh_from_db()
        self.v.approved_at=timezone.make_aware(datetime(2026,1,1));self.v.save(update_fields=['approved_at'])

    def test_owner_scope(self):
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(self.path).status_code,404)
        self.assertEqual(self.client.get('/api/schedules').data,[])

    def test_approval_only_assigned_director(self):
        self.client.post(self.path+'/submit',{'approver_id':self.users['technical'].pk},format='json')
        self.client.force_authenticate(self.users['finance'])
        self.assertEqual(self.client.post(self.path+'/approve',{},format='json').status_code,400)
        self.v.refresh_from_db();self.assertEqual(self.v.status,'review')

    def test_immutable_approved_and_new_revision(self):
        self.approve();old=deepcopy(self.v.data)
        self.client.force_authenticate(self.users['gip'])
        self.assertEqual(self.client.put(self.path,{'data':old,'edit_version':1},format='json').status_code,400)
        r=self.client.post(self.path+'/revise',{'reason':'Новые исходные данные'},format='json')
        self.assertEqual(r.status_code,200);self.assertEqual(r.data['revision']['number'],2)
        self.v.refresh_from_db();self.assertEqual(self.v.data,old)
        report=self.client.get('/api/reports?as_of=2026-01-20').data
        self.assertEqual(report[0]['revision']['number'],1)

    def test_conflict(self):
        r=self.client.put(self.path,{'data':self.v.data,'edit_version':999},format='json')
        self.assertEqual(r.status_code,412)

    def test_cost_mismatch_blocks_submission(self):
        self.v.data['budget']='101';self.v.save()
        r=self.client.post(self.path+'/submit',{'approver_id':self.users['technical'].pk},format='json')
        self.assertEqual(r.status_code,400)

    def test_fact_validations_and_permission(self):
        self.approve()
        payload={'node_id':'a','effective_date':'2026-01-20','progress':100,'actual_start':'2026-01-05'}
        self.assertEqual(self.client.post(self.path+'/facts',payload,format='json').status_code,403)
        self.client.force_authenticate(self.users['coordinator'])
        self.assertEqual(self.client.post(self.path+'/facts',payload,format='json').status_code,400)
        payload['actual_end']='2026-01-07'
        self.assertEqual(self.client.post(self.path+'/facts',payload,format='json').status_code,200)
        payload['actual_start']='2026-02-01'
        self.assertEqual(self.client.post(self.path+'/facts',payload,format='json').status_code,400)

    def test_history_known_then_vs_now(self):
        self.approve();self.client.force_authenticate(self.users['coordinator'])
        f=Fact.objects.create(schedule=self.s,node_id='b',progress=100,actual_start='2026-01-08',actual_end='2026-01-12',effective_date='2026-01-12',author=self.users['coordinator'])
        Fact.objects.filter(pk=f.pk).update(recorded_at=timezone.make_aware(datetime(2026,1,15,9)))
        now=self.client.get('/api/reports?as_of=2026-01-13').data[0]
        then=self.client.get('/api/reports?as_of=2026-01-13&known_at=2026-01-13').data[0]
        self.assertEqual(now['facts']['b']['progress'],100);self.assertNotIn('b',then['facts'])

    def test_zero_cost_work_prevents_completion(self):
        self.approve()
        for node in ['b','c']:
            Fact.objects.create(schedule=self.s,node_id=node,progress=100,effective_date='2026-01-20',actual_start='2026-01-08',actual_end='2026-01-14',author=self.users['coordinator'])
        r=self.client.get('/api/reports?as_of=2026-01-20').data[0]
        self.assertEqual(r['progress'],99.9);self.assertNotEqual(r['state'],'completed')

    def test_department_cost_masking(self):
        self.v.data['tasks'][1]['department']='Чужой отдел';self.v.save()
        self.client.force_authenticate(self.users['department'])
        r=self.client.get(self.path).data
        self.assertIsNone(r['data']['tasks'][1]['cost']);self.assertIsNone(r['baseline'])
        self.assertFalse(any(t['department']=='Чужой отдел' for t in r['load']))
        self.assertEqual(sum(float(t['cost']) for t in r['load']),40)

    def test_excel_codes_and_no_formula_execution(self):
        self.v.data['tasks'][0]['title']='=1+1';self.v.save()
        r=self.client.get(self.path+'/export')
        self.assertEqual(r.status_code,200)
        wb=load_workbook(BytesIO(r.content));sheet=wb['График А4']
        self.assertEqual(sheet['A8'].value,'062');self.assertEqual(sheet['B8'].data_type,'s');self.assertEqual(sheet['B8'].value,'=1+1')
        self.assertEqual(sheet.page_setup.orientation,'landscape')

    def test_snapshot_freezes_values(self):
        self.approve();self.client.force_authenticate(self.users['gip'])
        from .models import ReportSnapshot
        r=self.client.post('/api/snapshots?as_of=2026-01-20',{},format='json')
        self.assertEqual(r.status_code,201,r.data)
        saved=ReportSnapshot.objects.get(pk=r.data['id']);old=deepcopy(saved.data)
        Fact.objects.create(schedule=self.s,node_id='b',progress=50,effective_date='2026-01-12',actual_start='2026-01-08',author=self.users['coordinator'])
        saved.refresh_from_db();self.assertEqual(saved.data,old)
        self.assertEqual(self.client.get(f'/api/snapshots/{saved.pk}/export').status_code,200)

    def test_login_requires_csrf(self):
        c=Client(enforce_csrf_checks=True)
        self.assertEqual(c.post('/api/login',{'username':'gip','password':'TestPassword2026!'},content_type='application/json').status_code,403)

    def test_database_rejects_frozen_content_changes(self):
        from django.db import connection, DatabaseError, transaction
        if connection.vendor != 'postgresql':
            self.skipTest('PostgreSQL trigger')
        self.approve()
        changed = deepcopy(self.v.data)
        changed['budget'] = '999'
        with self.assertRaises(DatabaseError), transaction.atomic():
            Revision.objects.filter(pk=self.v.pk).update(data=changed)
        self.v.refresh_from_db()
        self.assertEqual(self.v.data['budget'], '100')

    def test_import_preview_not_applied(self):
        wb=Workbook();ws=wb.active
        ws.append(['ID','Работа','Шифр','Отдел','Автор','Начало','Окончание','Сумма'])
        ws.append(['new','New work','062','Путь','Автор','2026-01-05','2026-01-06',100])
        b=BytesIO();wb.save(b);b.seek(0);b.name='import.xlsx'
        r=self.client.post(self.path+'/import',{'file':b},format='multipart')
        self.assertEqual(r.status_code,200,r.data);self.assertEqual(r.data['tasks'][0]['specialty'],'062')
        self.v.refresh_from_db();self.assertEqual(len(self.v.data['tasks']),4)
